"""Post-process the generated deck.

The pptx_create spec cannot set font size on shapes, and the architecture slide
uses narrow boxes (~2 inches) where 18 pt text overflows. This sets explicit
sizes, small margins and middle anchoring so the diagram reads cleanly.
"""
from __future__ import annotations

import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

PATH = sys.argv[1] if len(sys.argv) > 1 else "docs/Aegis-SOC-Slides.pptx"
prs = Presentation(PATH)

# This script was written for the 10-slide generated deck. The final deck has 12
# slides and its own styling, so refuse rather than silently shrinking the fonts
# on a slide that is already correctly laid out.
if len(prs.slides) != 10:
    raise SystemExit(
        f"refusing: {PATH} has {len(prs.slides)} slides, this script only applies "
        "to the 10-slide generated deck (see scripts/finalize_deck.py for the final deck)"
    )

changed = 0
for idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        tf = shape.text_frame  # type: ignore[attr-defined]  # narrowed by has_text_frame
        is_diagram = idx == 3          # architecture slide
        if not is_diagram:
            continue
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(3)
        tf.margin_top = tf.margin_bottom = Pt(2)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size = Pt(11)
        changed += 1

prs.save(PATH)
print(f"styled {changed} shapes on the architecture slide -> {PATH}")
