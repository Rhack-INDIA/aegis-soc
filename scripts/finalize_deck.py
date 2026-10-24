"""Finalise the presentation deck: correct metadata, add speaker notes.

The deck's content and figures were already verified against the measured results;
this only fixes what a text dump cannot show:
  * lastModifiedBy was 'Steve Canny' and comments said 'generated using python-pptx'
    (both visible in PowerPoint File > Info)
  * no slide had speaker notes

    python scripts/finalize_deck.py IN.pptx OUT.pptx
"""
from __future__ import annotations

import sys

from pptx import Presentation

AUTHOR = "Rudra Pratap Dalei"

# One entry per slide, in order. Written to be spoken, not read.
NOTES = [
    "Open with the problem, not the product: a SOC has four jobs — collect, detect, triage, "
    "respond — and analyst time is the bottleneck. Then one line on what this is: that full loop "
    "in a single container, every detection mapped to ATT&CK, with an LLM in the triage seat. "
    "Finish with the hook: the model that reads logs is itself attackable, and I measured it.",

    "Three pressures drive the design question: alert volume grows faster than headcount, SIEM "
    "pricing scales with ingest rather than risk, and detection without triage and a safe response "
    "is only half a loop. Land the question — what is a credible SOC under a very small cloud "
    "budget — because every decision later in the deck traces back to it.",

    "Walk the five stages quickly, then stop on the safety principle: the model writes triage "
    "prose; deterministic code authorises containment. That single sentence is the thesis, and "
    "slide 7 is only survivable because of it.",

    "Name a decision at each component. Ingest: an 11-field ECS-lite subset, so it is not "
    "bespoke. Detection: two arms that fail differently — precise rules, and a baseline for "
    "behaviour with no signature. Triage: schema-validated JSON. Policy gate: severity, score, "
    "rule match, source IP. Store: SQLite locally, Firestore for the stateless Cloud Run design. "
    "If asked why one container: three services means three deploy paths and three fat images "
    "against a 0.5 GB registry cap.",

    "This is the live console. Point at the queue, then at the triage panel: verdict, confidence, "
    "narrative, and the audit footer that shows model mode and whether the schema validated. Say "
    "the queue is twelve alerts across eleven techniques — one row per detection, not per log "
    "line. If you are demonstrating, this is the moment to press Run attack simulation.",

    "Rules give precision and explainability; the baseline catches behaviour with no signature. "
    "The honest anecdote: my first fusion emitted only the highest-severity match, so a generic "
    "curl-pipe-bash rule masked the cron rule and an entire ATT&CK technique silently vanished "
    "from coverage. Only an ATT&CK-mapped evaluation could have caught that.",

    "Every field in a log line is written by the attacker. Twelve payloads asked the model to "
    "classify a HIGH-severity detection as benign; seven succeeded undefended. Five defences "
    "bring it to zero, and D4 — action separation — means even a successful hijack cannot trigger "
    "a block. Say why it matters: a silent downgrade is worse than a crash, because containment "
    "never fires and nobody notices.",

    "Lead with 1 out of 6, not 13 out of 13. The recall figure is circular — I wrote the corpus "
    "and the rules together — so it proves the pipeline is wired, nothing more. Six evasions: "
    "openssl instead of base64, systemd timers instead of cron, in-place truncation instead of "
    "rm. I attribute one. That is the number that describes the detector.",

    "Be exact here. The $0 figure is the Cloud Run reference design; this demonstration instance "
    "is EC2 at roughly $7.60 a month, because the organisation's GCP billing account was closed "
    "and Cloud Run requires active billing. Same image, same code, different host. Volunteering "
    "this is far stronger than being caught on it.",

    "Six gates run on every push. The image scan fails only on fixable findings — 44 unfixable OS "
    "CVEs are reported, not hidden, because a gate that always fails gets ignored. Seven findings "
    "fixed including three HIGH CVEs, four accepted with stated residual risk, zero open. Worth "
    "adding: the pipeline's own first runs failed on pipeline bugs, and I fixed the pipeline "
    "rather than the badge.",

    "Say all of these before anyone asks. Synthetic corpus and a single host, mock triage unless "
    "an API key is configured, simulated containment, no authentication over plain HTTP. An "
    "examiner who finds an undisclosed limitation marks down; one who hears it first marks up.",

    "Close on the constraint rather than the feature list: the free tier did not stop the project, "
    "it decided the architecture — no managed SQL produced the storage choice, a 0.5 GB image cap "
    "produced the statistical detector instead of a neural one, and no always-on VM produced the "
    "single-service design. Then invite questions.",
]


def main() -> None:
    src, dst = sys.argv[1], sys.argv[2]
    prs = Presentation(src)

    cp = prs.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.comments = ""
    cp.title = "Aegis-SOC — Professional Project Presentation"
    cp.subject = "Advanced Cloud Computing"

    for idx, slide in enumerate(prs.slides):
        if idx < len(NOTES):
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                tf.text = NOTES[idx]

    prs.save(dst)
    noted = sum(1 for s in prs.slides
                if (s.notes_slide.notes_text_frame is not None
                    and s.notes_slide.notes_text_frame.text.strip()))
    print(f"slides={len(prs.slides)} notes={noted} author={cp.author!r} -> {dst}")


if __name__ == "__main__":
    main()
